"""The QC log and the shot tracker (M4.3, QC_RULES "QC log structure").

The two sheets are tested from opposite directions. The QC log is ours, so the tests
ask whether it says what happened. The tracker is the studio's, so they mostly ask what
it does **not** write: 30 of its 39 columns belong to the vendor's team and a pasted row
that filled one would overwrite a fortnight of somebody else's tracking.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from proingest.core import exports, qc
from proingest.core.models import (
    AudioInfo,
    Batch,
    Deliverable,
    FrameRate,
    InOut,
    MediaInfo,
    QCResult,
    ShotRow,
    Turnover,
)
from tests.fixtures.names import identity_of

RATE_24 = FrameRate(24)
RATE_2398 = FrameRate(24000, 1001)


def deliverable(kind: str, name: str, res: str | None = None, **kwargs: object) -> Deliverable:
    return Deliverable(
        kind=kind,
        name=name,
        path=Path("/delivery/MELT/MELT0001") / name,
        version=1,
        res=res,
        status=kwargs.pop("status", "done"),  # type: ignore[arg-type]
        **kwargs,  # type: ignore[arg-type]
    )


def row(clip_name: str = "MELT0001_pl01", rate: FrameRate = RATE_24, **kwargs: object) -> ShotRow:
    """A delivered plate: 4k and HD sequences, both references, audio and side files."""
    built = ShotRow(
        turnover_id="turnover001",
        clip_name=clip_name,
        identity=identity_of(clip_name),
        media=MediaInfo(
            path=Path("/turnover/MELT0001_pl01.mov"),
            codec="prores",
            pixel_format="yuv444p12le",
            width=3840,
            height=2160,
            rate=rate,
            frame_count=240,
            start_frame=0,
            start_timecode=86400,
        ),
        snapshot=InOut(8, 231),
        current=InOut(8, 231),
        audio_path=Path("/turnover/MELT0001_pl01.wav"),
        audio=AudioInfo(
            path=Path("/turnover/MELT0001_pl01.wav"),
            duration_samples=448000,
            sample_rate=48000,
            channels=2,
            bit_depth=24,
        ),
    )
    built.deliverables = [
        deliverable("raw_dir", "MELT0001_pl01_raw_4k_v01", "4k", frame_count=224),
        deliverable("raw_dir", "MELT0001_pl01_raw_HD_v01", "HD", frame_count=224),
        deliverable("ref_mp4", "MELT0001_pl01_ref_4k_v01.mp4", "4k", checksum="aa"),
        deliverable("ref_mp4", "MELT0001_pl01_ref_HD_v01.mp4", "HD", checksum="bb"),
        deliverable("audio", "MELT0001_pl01_audio_v01.wav", checksum="cc"),
        deliverable("hdri", "MELT0001_pl01_HDRI_v01.exr", checksum="dd"),
        deliverable("camdata", "MELT0001_pl01_camData_v01.txt", checksum="ee"),
    ]
    for key, value in kwargs.items():
        setattr(built, key, value)
    return built


def batch_of(*rows: ShotRow, name: str = "melt") -> Batch:
    return Batch(
        name=name,
        delivery_root=Path("/delivery"),
        turnovers=[Turnover(turnover_id="turnover001", folder=Path("/turnover"))],
        rows=list(rows),
    )


def sheet_rows(path: Path, title: str) -> list[list[object]]:
    return [list(values) for values in load_workbook(path)[title].values]


class TestReportPaths:
    def test_both_land_in_the_show_s_reports_folder(self) -> None:
        log, tracker = exports.report_paths(batch_of(row()), Path("/d"), date(2026, 9, 11))
        assert log.parent == Path("/d/MELT/_reports")
        assert tracker.parent == log.parent

    def test_the_names_carry_the_batch_and_the_date(self) -> None:
        log, tracker = exports.report_paths(batch_of(row()), Path("/d"), date(2026, 9, 11))
        assert log.name == "qc_ingest_log_melt_20260911.xlsx"
        assert tracker.name == "shot_tracker_melt_20260911.xlsx"

    def test_a_batch_with_no_shot_code_says_so(self) -> None:
        """There is no show to file under, and guessing one writes to the wrong place."""
        with pytest.raises(ValueError, match="no show"):
            exports.report_paths(batch_of(ShotRow(turnover_id="t", clip_name="junk")), Path("/d"))


class TestTheWrite:
    def test_both_files_land_atomically(self, tmp_path: Path) -> None:
        """Written to a temp name and renamed, like every deliverable: a crash mid-save
        must not leave a file that looks finished, and no temp name survives."""
        exports.write_qc_log(batch_of(row()), tmp_path / "r" / "log.xlsx")
        exports.write_shot_tracker(batch_of(row()), tmp_path / "r" / "tracker.xlsx")
        assert sorted(p.name for p in (tmp_path / "r").iterdir()) == ["log.xlsx", "tracker.xlsx"]

    def test_a_control_character_in_a_note_is_dropped_not_raised(self, tmp_path: Path) -> None:
        """F15: openpyxl refuses one, and the raise used to leave Run greyed for good."""
        noted = row(notes="sky \x07replace")
        written = exports.write_qc_log(batch_of(noted), tmp_path / "log.xlsx")
        assert written.is_file()


class TestQcLogSheets:
    @pytest.fixture
    def log(self, tmp_path: Path) -> Path:
        return exports.write_qc_log(batch_of(row()), tmp_path / "log.xlsx", date(2026, 9, 11))

    def test_the_summary_records_the_ffmpeg_version(self, log: Path) -> None:
        """PACKAGING.md: every QC log names the ffmpeg that made the deliverables."""
        summary = {line[0]: line[1] for line in sheet_rows(log, "Summary")}
        assert "ffmpeg version" in str(summary["ffmpeg"]).lower()

    def test_the_three_sheets_the_spec_names(self, log: Path) -> None:
        """Three since 2026-09-22: Side Files and Camera Data went with the deliverables
        they described, which the tool no longer produces."""
        assert load_workbook(log).sheetnames == ["Summary", "Shots", "Deliverables"]

    def test_the_summary_counts_what_ran(self, log: Path) -> None:
        summary: dict[object, object] = {key: value for key, value in sheet_rows(log, "Summary")[1:]}
        assert summary["Batch"] == "melt"
        assert summary["Rows"] == 1
        assert summary["Rows delivered"] == 1
        assert summary["Deliverables"] == 7

    def test_one_shots_row_per_shot(self, log: Path) -> None:
        rows = sheet_rows(log, "Shots")
        assert rows[0][:4] == ["Turnover", "Clip name", "Shot code", "Elem"]
        assert len(rows) == 2
        assert rows[1][:4] == ["turnover001", "MELT0001_pl01", "MELT0001", "pl01"]

    def test_the_shots_sheet_carries_timecode_as_well_as_frames(self, log: Path) -> None:
        """Frames are what the tool computes in; timecode is what a human reads."""
        header, values = sheet_rows(log, "Shots")
        row_of = dict(zip(header, values, strict=True))
        assert row_of["Delivered In"] == 8
        assert row_of["Delivered In TC"] == "01:00:00:08"

    def test_the_shots_sheet_names_the_encoding_the_clip_asked_for(self, tmp_path: Path) -> None:
        """Verbatim, because what QC-047 needs corrected is the string somebody typed."""
        named = row()
        named.source_encoding = "C-Log3"
        path = tmp_path / "log.xlsx"
        exports.write_qc_log(Batch(name="b", rows=[named]), path)
        header, values = sheet_rows(path, "Shots")
        assert dict(zip(header, values, strict=True))["Source encoding"] == "C-Log3"

    def test_a_clip_that_named_no_encoding_leaves_the_column_empty(self, log: Path) -> None:
        header, values = sheet_rows(log, "Shots")
        assert not dict(zip(header, values, strict=True))["Source encoding"]

    def test_one_deliverables_row_each(self, log: Path) -> None:
        assert len(sheet_rows(log, "Deliverables")) == 8


class TestDeliverableRuleColumns:
    """One column per QC-1xx rule, which is what makes an NA mean anything."""

    def test_a_column_for_every_deliverable_rule(self, tmp_path: Path) -> None:
        log = exports.write_qc_log(batch_of(row()), tmp_path / "log.xlsx")
        header = sheet_rows(log, "Deliverables")[0]
        assert header[-len(qc.DELIVERABLE_RULES) :] == list(qc.DELIVERABLE_RULES)

    def test_a_rule_a_kind_does_not_owe_reads_na(self, tmp_path: Path) -> None:
        """An mp4 has no opinion about EXR compression, so QC-105 is NA and not PASS."""
        log = exports.write_qc_log(batch_of(row()), tmp_path / "log.xlsx")
        header, *rows = sheet_rows(log, "Deliverables")
        by_name = {
            str(values[5]).rsplit("/", 1)[-1]: dict(zip(header, values, strict=True)) for values in rows
        }
        assert by_name["MELT0001_pl01_ref_HD_v01.mp4"]["QC-105"] == "NA"
        assert by_name["MELT0001_pl01_raw_4k_v01"]["QC-105"] == "PASS"

    def test_a_recorded_failure_reads_fail(self, tmp_path: Path) -> None:
        target = row()
        target.deliverables[0].qc.append(
            QCResult("QC-102", "error", "deliverable", "frame numbering has a gap")
        )
        log = exports.write_qc_log(batch_of(target), tmp_path / "log.xlsx")
        header, *rows = sheet_rows(log, "Deliverables")
        assert dict(zip(header, rows[0], strict=True))["QC-102"] == "FAIL"

    def test_a_render_that_never_happened_is_na_for_everything_else(self, tmp_path: Path) -> None:
        target = row()
        target.deliverables[0].status = "failed"
        target.deliverables[0].qc.append(QCResult("QC-100", "error", "deliverable", "no file"))
        log = exports.write_qc_log(batch_of(target), tmp_path / "log.xlsx")
        header, *rows = sheet_rows(log, "Deliverables")
        first = dict(zip(header, rows[0], strict=True))
        assert first["QC-100"] == "FAIL"
        assert first["QC-101"] == "NA"


class TestShotTracker:
    @pytest.fixture
    def tracker(self, tmp_path: Path) -> Path:
        return exports.write_shot_tracker(batch_of(row()), tmp_path / "tracker.xlsx")

    def test_the_header_is_the_studio_s_own(self, tracker: Path) -> None:
        header = sheet_rows(tracker, "Shots")[0]
        assert len(header) == 39
        assert header[2] == "Shot Code (ABCD123)"
        assert header[34] == "Turnover Stringout (Edit)"

    def test_the_columns_the_tool_owns(self, tracker: Path) -> None:
        values = sheet_rows(tracker, "Shots")[1]
        assert values[2] == "MELT0001"
        assert values[3] == "MELT0001"
        assert values[4] == "MELT0001_pl01_ref_HD_v01.mp4"
        assert values[5] is None, "HDRI is the studio's column and no longer ours to fill"
        assert values[6] is None, "CAM Data likewise"
        assert values[8] == "24"
        assert values[9] == "✓"

    def test_every_other_column_is_empty(self, tracker: Path) -> None:
        """The whole point: 30 columns belong to the vendor and a paste must not touch them."""
        values = sheet_rows(tracker, "Shots")[1]
        for index, value in enumerate(values):
            if index not in exports.TOOL_OWNED_COLUMNS and index != 0:
                assert value is None, f"column {index} ({exports.TRACKER_HEADERS[index]}) was written"

    def test_the_plate_video_is_the_hd_reference_not_the_4k(self, tracker: Path) -> None:
        """Every one of the 782 real values in the studio sheet is the HD one."""
        assert sheet_rows(tracker, "Shots")[1][4] == "MELT0001_pl01_ref_HD_v01.mp4"

    def test_the_plate_marks_read_like_the_sheet_s_own(self, tracker: Path) -> None:
        assert sheet_rows(tracker, "Shots")[1][7] == "4K ✓\nHD ✓"

    def test_a_missing_plate_is_marked_rather_than_blank(self, tmp_path: Path) -> None:
        target = row()
        target.deliverables = [d for d in target.deliverables if d.res != "HD"]
        written = exports.write_shot_tracker(batch_of(target), tmp_path / "t.xlsx")
        assert sheet_rows(written, "Shots")[1][7] == "4K ✓\nHD \u2014"

    def test_the_stringout_column_is_left_for_a_human(self, tracker: Path) -> None:
        """OQ-41: the grammar the tool would rebuild it from matches none of the real names."""
        assert sheet_rows(tracker, "Shots")[1][34] is None

    def test_a_23976_source_is_recorded_at_the_24_it_was_delivered_at(self, tmp_path: Path) -> None:
        """Every deliverable is written at 24 frame for frame (user, 2026-09-23)."""
        written = exports.write_shot_tracker(batch_of(row(rate=RATE_2398)), tmp_path / "t.xlsx")
        assert sheet_rows(written, "Shots")[1][8] == "24"

    def test_one_line_per_shot_code(self, tmp_path: Path) -> None:
        """The studio sheet is one line per shot; a shot's plates and stills are one line."""
        clean = row(clip_name="MELT0001_cp01")
        clean.deliverables = [deliverable("raw_dir", "MELT0001_cp01_raw_4k_v01", "4k")]
        other = row(clip_name="MELT0002_pl01")
        written = exports.write_shot_tracker(batch_of(clean, row(), other), tmp_path / "t.xlsx")
        lines = sheet_rows(written, "Shots")[1:]
        assert [line[2] for line in lines] == ["MELT0001", "MELT0002"]
        assert lines[0][4] == "MELT0001_pl01_ref_HD_v01.mp4", "described by its pl, not its cp"

    def test_a_shot_with_no_plate_has_no_plate_video_and_no_marks(self, tmp_path: Path) -> None:
        clean = row(clip_name="MELT0001_cp01")
        written = exports.write_shot_tracker(batch_of(clean), tmp_path / "t.xlsx")
        line = sheet_rows(written, "Shots")[1]
        assert line[4] is None
        assert line[7] == "4K \u2014\nHD \u2014"
        assert line[9] is None

    def test_a_blocked_row_is_not_offered_for_pasting(self, tmp_path: Path) -> None:
        blocked = row()
        blocked.qc.append(QCResult("QC-066", "error", "row", "no event"))
        written = exports.write_shot_tracker(batch_of(blocked), tmp_path / "t.xlsx")
        assert len(sheet_rows(written, "Shots")) == 1

    def test_a_cancelled_or_failed_render_is_not_offered_for_pasting(self, tmp_path: Path) -> None:
        cancelled = row()
        cancelled.deliverables[0].status = "skipped"
        failed = row(clip_name="MELT0002_pl01")
        failed.deliverables[-1].status = "failed"
        written = exports.write_shot_tracker(batch_of(cancelled, failed), tmp_path / "t.xlsx")
        assert len(sheet_rows(written, "Shots")) == 1

    def test_a_skipped_row_is_not_offered_for_pasting(self, tmp_path: Path) -> None:
        """It delivered nothing, so it is not a shot in the delivery."""
        written = exports.write_shot_tracker(
            batch_of(row(), row(clip_name="MELT0002_pl01", skipped=True)), tmp_path / "t.xlsx"
        )
        assert len(sheet_rows(written, "Shots")) == 2
